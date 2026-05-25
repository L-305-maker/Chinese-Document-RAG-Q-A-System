from pathlib import Path
from typing import List
from abc import ABC,abstractmethod

from docx import Document as DocxDocument
from pypdf import PdfReader

from src.document import Document


class Baseloader():

    #模板
    
    @abstractmethod
    def load(self,file_Path:str)->List[Document]:
        pass

class pdf_loader(Baseloader):
    def load(self,file_Path:str)->List[Document]:
        path = Path(file_Path)
        reader = PdfReader(file_Path)

        documents = []

        for page_index,page in enumerate(reader.pages):
            text = page.extract_text() or ""
            documents.append(Document(
                page_content = text,
                    metadata = {
                        "source":path.name,
                        "file_path":str(file_Path),
                        "file_type":"pdf",
                        "page":page_index+1
                    }
            ))

        return documents
    
class word_loader(Baseloader):
    def load(self,file_Path:str)->List[Document]:
        path = Path(file_Path)
        docx = DocxDocument(file_Path)

        paragraphs = []

        for paragraph in docx.paragraphs:
             text = paragraph.text.strip()
             if text:
                 paragraphs.append(text)
        
        full_text = "\n".join(paragraphs)

        documents = []
        documents.append(Document(
            page_content = full_text,
            metadata = {
                "source":path.name,
                "file_path":str(file_Path),
                "file_type":"word"
            }
        ))

        return documents
    
class Txtloader(Baseloader):
    def load(self,file_Path:str)->List[Document]:
        path = Path(file_Path)
        with open(file_Path, "r", encoding="utf-8") as f:
            text = f.read()

        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_path": str(path),
                    "file_type": "txt",
                },
            )
        ]

class Markdownloader(Baseloader):
    def load(self,file_Path:str)->List[Document]:
        path = Path(file_Path)

        with open(file_Path, "r", encoding="utf-8") as f:
            text = f.read()

        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_path": str(path),
                    "file_type": "markdown",
                },
            )
        ]


#将各个文件统一输入到LoaderFactory进行处理
class LoaderFactory():
    def __init__(self):
        self.loaders = {
            "txt": Txtloader(),
            "markdown": Markdownloader(),
            "pdf": pdf_loader(),
            "word": word_loader(),
        }

    def get_loader(self, file_type: str) -> Baseloader:
        if file_type not in self.loaders:
            raise ValueError(f"No loader found for file type: {file_type}")
        return self.loaders[file_type]